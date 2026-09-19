#!/usr/bin/env python3
"""E2E + unit test for SHELL + COMPONENT composition (both modes).

Self-contained (writes its own fixtures to working/e2e-test-shell). Asserts:

  [1] the shells/components build cleanly and exist;
  [2] mount unit — a component's shapes are translated to the region origin and
      all cNvPr ids stay unique after mounting (the one corruption the mount
      engine could introduce);
  [3] capacity guard — an over-capacity component emits a WARN;
  [4] freestyle e2e — a deck with a `halves` composite slide composes, qa is
      clean, and the shell's content (stats + card bullets) survives into the deck;
  [5] proposal e2e + regression — composing a real proposal plan with one
      generated slide authored as a shell changes ONLY that slide; the rest of
      the deck is byte-identical to the non-shell compose.

Like the other e2e tests, it does NOT exercise the LLM steps — the digest/plan/
content JSONs are written deterministically here.

Usage: python scripts/setup/run_e2e_test_shell.py
Exit codes: 0 = pass, 1 = fail.
"""

import contextlib
import filecmp
import io
import json
import subprocess
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parent.parent          # .../pptx-editor/scripts
SKILL = SCRIPTS.parent                                     # .../pptx-editor
REPO = SCRIPTS.parents[4]                                  # repo root
RES = SKILL / "resources"
SHELLS = RES / "shells"
COMPONENTS = RES / "components"
TEMPLATE_DIR = RES / "2026_ProposalsAI_Master_Template"
MANIFEST = TEMPLATE_DIR / "2026_ProposalsAI_Master_Template-manifest.json"
OUT = REPO / "working/e2e-test-shell"

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
EMU = 914400

sys.path.insert(0, str(SCRIPTS))


def run(args, **kw):
    """Run one of our own scripts via sys.executable (no shell, no external binaries)."""
    print("  $", " ".join(str(a) for a in args))
    return subprocess.run([sys.executable, *map(str, args)], shell=False, **kw)


def fail(msg):
    print(f"FAIL — {msg}")
    return 1


def deck_text(pptx: Path) -> str:
    from pptx import Presentation
    out = []
    for s in Presentation(str(pptx)).slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
    return "\n".join(out)


def _ids(root) -> list:
    return [int(c.get("id")) for c in root.iter(f"{P}cNvPr") if c.get("id", "").isdigit()]


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)

    # --- [1] build pieces ---------------------------------------------------
    print("[1/5] build shells + components")
    if run([SCRIPTS / "setup/build_components.py"]).returncode:
        return fail("build_components.py")
    if run([SCRIPTS / "setup/build_shells.py"]).returncode:
        return fail("build_shells.py")
    for d in ("shells/halves", "components/card", "components/stat-stack", "components/mini-table"):
        if not (RES / d / "slots.json").is_file():
            return fail(f"missing built piece {d}")
    print("  ok — halves + card/stat-stack/mini-table built")

    from fill_blueprint import fill_blueprint
    from mount import mount_component, max_shape_id

    # --- [2] mount unit: translation + id uniqueness ------------------------
    print("[2/5] mount unit — translation + id uniqueness")
    meta = json.loads((SHELLS / "halves" / "slots.json").read_text(encoding="utf-8"))
    region = next(r for r in meta["regions"] if r["name"] == "right")
    root = ET.fromstring(fill_blueprint(SHELLS / "halves", {"SECTION": "S", "TITLE": "T"}))
    spt = root.find(f"{P}cSld/{P}spTree")
    block = {"__component__": "stat-stack", "STAT1_VALUE": "42%", "STAT1_LABEL": "X"}
    mount_component(root, region, COMPONENTS / "stat-stack", block, max_shape_id(spt) + 1, OUT, 9)
    region_x = int(round(region["x"] * EMU))
    val_off = None
    for sp in spt:
        if any(t.text == "42%" for t in sp.iter(f"{A}t")):
            off = sp.find(f"{P}spPr/{A}xfrm/{A}off")
            val_off = int(off.get("x"))
    if val_off is None:
        return fail("mounted stat value '42%' not found in shell")
    if abs(val_off - region_x) > 2:
        return fail(f"translation wrong: stat off.x={val_off}, region.x={region_x}")
    ids = _ids(root)
    if len(ids) != len(set(ids)):
        return fail("duplicate cNvPr ids after mount")
    print(f"  ok — translated to region origin (x={val_off}), {len(ids)} ids all unique")

    # --- [3] capacity guard -------------------------------------------------
    print("[3/5] capacity guard — over-capacity component WARNs")
    r2 = ET.fromstring(fill_blueprint(SHELLS / "halves", {"SECTION": "S", "TITLE": "T"}))
    left = next(r for r in meta["regions"] if r["name"] == "left")
    over = {"__component__": "card", "CARD_HEADER": "H",
            "CARD_BODY": ["a", "b", "c", "d", "e", "f", "g"]}  # 7 > cap 5
    buf = io.StringIO()
    with contextlib.redirect_stderr(buf):
        mount_component(r2, left, COMPONENTS / "card", over,
                        max_shape_id(r2.find(f"{P}cSld/{P}spTree")) + 1, OUT, 9)
    if "WARN" not in buf.getvalue():
        return fail("no capacity WARN for a 7-bullet card (cap 5)")
    print("  ok — capacity WARN emitted on overflow")

    # --- [4] freestyle e2e: a composite shell slide -------------------------
    print("[4/5] freestyle e2e — halves shell in a deck")
    free = OUT / "free"
    free.mkdir(parents=True, exist_ok=True)
    (free / "plan.json").write_text(json.dumps({"mode": "free", "slides": [
        {"n": 1, "type": "generated", "blueprint": "three-col-cards"},
        {"n": 2, "type": "generated", "shell": "halves"},
        {"n": 3, "type": "generated", "blueprint": "numbered-pipeline"},
    ]}), encoding="utf-8")
    (free / "content.json").write_text(json.dumps({"placeholders": {}, "generated": {
        "1": {"SECTION": "A", "TITLE": "Pilares", "CARD1_HEADER": "Uno", "CARD1_BODY": ["x"],
              "CARD2_HEADER": "Dos", "CARD2_BODY": ["y"], "CARD3_HEADER": "Tres", "CARD3_BODY": ["z"]},
        "2": {"SECTION": "IMPACTO", "TITLE": "Metricas y entrega", "__shell__": "halves",
              "left": {"__component__": "stat-stack", "STAT1_VALUE": "55pct", "STAT1_LABEL": "MENOS COSTE",
                       "STAT2_VALUE": "3x", "STAT2_LABEL": "MAS RAPIDO"},
              "right": {"__component__": "card", "CARD_HEADER": "Que incluye",
                        "CARD_BODY": ["Arquitectura cloud", "Pipeline CI/CD", "Observabilidad"]}},
        "3": {"SECTION": "PLAN", "TITLE": "Fases", "STEP1_HEADER": "Discovery", "STEP1_BODY": "Alcance",
              "STEP2_HEADER": "Build", "STEP2_BODY": "Iterativo"},
    }}), encoding="utf-8")
    if run([SCRIPTS / "compose_free.py", free / "plan.json", free / "content.json",
            TEMPLATE_DIR, free, "--pack"]).returncode:
        return fail("compose_free.py (shell deck)")
    if run([SCRIPTS / "qa.py", free / "presentation.pptx"]).returncode:
        return fail("qa.py flagged the freestyle shell deck")
    text = deck_text(free / "presentation.pptx")
    for needle in ("55pct", "MENOS COSTE", "Que incluye", "Pipeline CI/CD"):
        if needle not in text:
            return fail(f"shell content '{needle}' missing from the deck")
    from pptx import Presentation
    if len(Presentation(str(free / "presentation.pptx")).slides) != 3:
        return fail("freestyle shell deck should have 3 slides")
    print("  ok — composite shell composed, qa clean, content survived")

    # --- [5] proposal e2e + regression --------------------------------------
    print("[5/5] proposal e2e — shell replaces one slide, rest unchanged")
    prop = OUT / "prop"
    prop.mkdir(parents=True, exist_ok=True)
    (prop / "digest.json").write_text(json.dumps({"collab": "fixed_price", "delivery": "plain_team",
        "flags": {"has_discovery": False, "phased": False, "is_brownfield": False,
                  "tech_lead_global": True}}), encoding="utf-8")
    if run([SCRIPTS / "plan.py", MANIFEST, prop / "digest.json", prop / "plan.json"]).returncode:
        return fail("plan.py")
    plan = json.loads((prop / "plan.json").read_text(encoding="utf-8"))
    gen = [s for s in plan["slides"] if s["type"] == "generated"]
    if not all("shell_options" in s for s in gen):
        return fail("plan.py did not emit shell_options on every generated slide")
    g = gen[0]["n"]  # author the first generated slide two ways
    base = {"placeholders": {}, "generated": {str(g): {"SECTION": "NUESTRA SOLUCION", "TITLE": "X"}}}
    shell = {"placeholders": {}, "generated": {str(g): {"SECTION": "NUESTRA SOLUCION", "TITLE": "X",
        "__shell__": "halves",
        "left": {"__component__": "stat-stack", "STAT1_VALUE": "30pct", "STAT1_LABEL": "TTM"},
        "right": {"__component__": "card", "CARD_HEADER": "Entregables", "CARD_BODY": ["Catalogo", "Checkout"]}}}}
    (prop / "c-base.json").write_text(json.dumps(base), encoding="utf-8")
    (prop / "c-shell.json").write_text(json.dumps(shell), encoding="utf-8")
    for tag, cfile in (("base", "c-base.json"), ("shell", "c-shell.json")):
        if run([SCRIPTS / "compose.py", prop / "plan.json", prop / cfile,
                TEMPLATE_DIR, prop / tag, "--pack"]).returncode:
            return fail(f"compose.py ({tag})")
    bdir = prop / "base/working-unpacked/ppt/slides"
    sdir = prop / "shell/working-unpacked/ppt/slides"
    base_slides = sorted(p.name for p in bdir.glob("slide*.xml"))
    shell_slides = sorted(p.name for p in sdir.glob("slide*.xml"))
    if base_slides != shell_slides:
        return fail("shell compose changed the deck's slide set")
    differing = [n for n in base_slides if not filecmp.cmp(bdir / n, sdir / n, shallow=False)]
    if differing != [f"slide{g}.xml"]:
        return fail(f"expected only slide{g}.xml to differ, got {differing}")
    # the shell slide itself must be clean
    raw = (sdir / f"slide{g}.xml").read_text(encoding="utf-8")
    if "___SLOT_" in raw:
        return fail(f"slide{g}.xml has unfilled markers")
    sroot = ET.fromstring(raw)
    sids = _ids(sroot)
    if len(sids) != len(set(sids)):
        return fail(f"slide{g}.xml has duplicate ids")
    print(f"  ok — only slide{g}.xml changed ({len(base_slides)} slides total), shell slide clean")

    print("\nPASS — shell+component E2E clean (mount unit + capacity + freestyle + proposal regression).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
