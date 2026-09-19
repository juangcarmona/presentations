#!/usr/bin/env python3
"""Deterministic, self-contained E2E regression test for the FREESTYLE pipeline.

Freestyle (non-proposal) counterpart of run_e2e_test.py and a focused sibling of
run_e2e_test_shell.py's freestyle section: it exercises the plain blueprint path
(no shells). It writes its OWN reference fixtures (digest / plan / content),
runs the Python half of the freestyle pipeline (compose_free -> pack -> qa) and
asserts:

  1. compose_free + pack succeed and qa.py reports a clean deck;
  2. the deck has exactly one slide per planned topic;
  3. no two consecutive slides share a blueprint (layout-variety hard rule);
  4. CONTENT COVERAGE — every topic heading from the source digest, plus the
     distinctive source terms, actually appear in the rendered slides. This is
     what proves the deck is *about the source content*, not just structurally OK.

Like the other e2e tests, it does NOT exercise the LLM steps (digest-free,
plan-free, outline, content) — those are replaced by the inline reference JSONs.

Usage:
    python scripts/setup/run_e2e_test_free.py
Exit codes: 0 = pass, 1 = fail.
"""

import json
import subprocess
import sys
import unicodedata
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parent.parent          # .../pptx-editor/scripts
SKILL = SCRIPTS.parent                                     # .../pptx-editor
REPO = SCRIPTS.parents[4]                                  # repo root
TEMPLATE_DIR = SKILL / "resources" / "2026_ProposalsAI_Master_Template"
OUT = REPO / "working" / "e2e-test-free"

# --- inline reference fixtures (the freestyle source, deterministic) ----------
# Four topics, one slide each, with four DISTINCT blueprints (no consecutive
# repeats). Mirrors what digest-free -> plan-free -> outline -> content emit.
DIGEST = {
    "mode": "free",
    "title": "Internal Developer Platform — 2026 Rollout",
    "language": "en",
    "audience": "Engineering leadership",
    "topics": [
        {"heading": "Three platform pillars", "kind": "parallel",
         "points": ["Golden paths", "Self-service", "Observability"], "notes": ""},
        {"heading": "How we roll it out", "kind": "process",
         "points": ["Discover", "Pilot", "Expand", "Optimize"], "notes": ""},
        {"heading": "What teams get on day one", "kind": "list",
         "points": ["Environments", "Pipelines", "Security", "Support"], "notes": ""},
        {"heading": "Measurable outcomes", "kind": "parallel",
         "points": ["Faster lead time", "More deployments", "Less setup", "High availability"],
         "notes": ""},
    ],
}

PLAN = {
    "mode": "free",
    "language": "en",
    "generated_count": 4,
    "slides": [
        {"n": 1, "type": "generated", "blueprint": "three-col-cards",
         "blueprint_options": ["three-col-cards", "benefits-icons", "summary-cards"],
         "section_label": "Three platform pillars", "group_label": "Three platform pillars"},
        {"n": 2, "type": "generated", "blueprint": "numbered-pipeline",
         "blueprint_options": ["numbered-pipeline", "phase-dark-left", "phase-detail-left-right"],
         "section_label": "How we roll it out", "group_label": "How we roll it out"},
        {"n": 3, "type": "generated", "blueprint": "icon-list-split",
         "blueprint_options": ["icon-list-split", "numbered-list-callout"],
         "section_label": "What teams get on day one", "group_label": "What teams get on day one"},
        {"n": 4, "type": "generated", "blueprint": "summary-cards",
         "blueprint_options": ["summary-cards", "three-col-cards", "benefits-icons"],
         "section_label": "Measurable outcomes", "group_label": "Measurable outcomes"},
    ],
}

CONTENT = {
    "placeholders": {},
    "generated": {
        "1": {
            "SECTION": "PLATFORM PILLARS",
            "TITLE": "Three pillars of the Developer Platform",
            "CARD1_HEADER": "Golden paths",
            "CARD1_BODY": ["Scaffolding templates", "Paved CI/CD", "Secure defaults"],
            "CARD2_HEADER": "Self-service",
            "CARD2_BODY": ["One-click environments", "Service catalog", "Automated provisioning"],
            "CARD3_HEADER": "Observability",
            "CARD3_BODY": ["Unified dashboards", "Distributed tracing", "Actionable alerts"],
        },
        "2": {
            "SECTION": "ROLLOUT",
            "TITLE": "How we roll it out",
            "STEP1_HEADER": "Discover", "STEP1_BODY": "Audit current tooling and pain points",
            "STEP2_HEADER": "Pilot", "STEP2_BODY": "Onboard two teams onto golden paths",
            "STEP3_HEADER": "Expand", "STEP3_BODY": "Scale to all product squads",
            "STEP4_HEADER": "Optimize", "STEP4_BODY": "Iterate on metrics and feedback",
        },
        "3": {
            "SECTION": "WHAT TEAMS GET",
            "TITLE": "What teams get on day one",
            "INTRO": "Everything a squad needs to ship safely.",
            "ROW1_LABEL": "Environments", "ROW1_VALUE": "Provisioned in minutes, not days",
            "ROW2_LABEL": "Pipelines", "ROW2_VALUE": "Pre-wired build, test and deploy",
            "ROW3_LABEL": "Security", "ROW3_VALUE": "Policy and secrets baked in",
            "ROW4_LABEL": "Support", "ROW4_VALUE": "Platform team on call 24/7",
        },
        "4": {
            "SECTION": "OUTCOMES",
            "TITLE": "Measurable outcomes",
            "CARD1_HEADER": "40%", "CARD1_BODY": "Faster lead time to production",
            "CARD2_HEADER": "3x", "CARD2_BODY": "More frequent deployments",
            "CARD3_HEADER": "60%", "CARD3_BODY": "Less time on environment setup",
            "CARD4_HEADER": "99.9%", "CARD4_BODY": "Platform availability",
        },
    },
}

# Distinctive content terms that must survive into the rendered deck.
REQUIRED_TERMS = ["Golden paths", "Self-service", "Observability", "Discover", "deployments"]


def run(args, **kw):
    """Run one of our own scripts via sys.executable (no shell, no external binaries)."""
    print("  $", " ".join(str(a) for a in args))
    return subprocess.run([sys.executable, *map(str, args)], shell=False, **kw)


def _norm(s: str) -> str:
    """Lowercase + strip accents so 'Qué' matches 'que' robustly."""
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s.lower()


def _deck_text(pptx: Path) -> str:
    from pptx import Presentation
    p = Presentation(str(pptx))
    chunks = []
    for slide in p.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
    return _norm("\n".join(chunks))


def main() -> int:
    if not TEMPLATE_DIR.exists():
        print(f"FAIL — missing template dir: {TEMPLATE_DIR}")
        return 1

    OUT.mkdir(parents=True, exist_ok=True)
    plan = OUT / "slide-plan.json"
    content = OUT / "content-blocks.json"
    plan.write_text(json.dumps(PLAN, ensure_ascii=False, indent=2), encoding="utf-8")
    content.write_text(json.dumps(CONTENT, ensure_ascii=False, indent=2), encoding="utf-8")
    pptx = OUT / "presentation.pptx"

    print("[1/4] compose_free.py --pack")
    r = run([SCRIPTS / "compose_free.py", plan, content, TEMPLATE_DIR, OUT, "--pack"])
    if r.returncode != 0:
        print("FAIL — compose_free.py")
        return 1

    print("[2/4] qa.py")
    r = run([SCRIPTS / "qa.py", pptx])
    if r.returncode != 0:
        print("FAIL — qa.py reported issues")
        return 1

    slides = PLAN["slides"]

    print("[3/4] structure — slide count + layout variety")
    from pptx import Presentation
    n_slides = len(Presentation(str(pptx)).slides)
    if n_slides != len(slides):
        print(f"FAIL — expected {len(slides)} slides, deck has {n_slides}")
        return 1
    for a, b in zip(slides, slides[1:]):
        if a["blueprint"] == b["blueprint"]:
            print(f"FAIL — consecutive slides share blueprint '{a['blueprint']}' "
                  f"(n={a['n']} & n={b['n']})")
            return 1
    print(f"  ok — {n_slides} slides, no consecutive layout repeats")

    print("[4/4] content coverage — slides are about the source content")
    text = _deck_text(pptx)
    missing = []
    for topic in DIGEST["topics"]:
        # heading should be recognisable on its slide; check its distinctive word
        head = _norm(topic["heading"])
        key = max(head.replace("(", " ").replace(")", " ").split(), key=len)
        if key not in text:
            missing.append(f"topic heading «{topic['heading']}» (key '{key}')")
    for term in REQUIRED_TERMS:
        if _norm(term) not in text:
            missing.append(f"term «{term}»")
    if missing:
        print("FAIL — the deck does not cover the source content:")
        for m in missing:
            print(f"    - missing {m}")
        return 1
    print(f"  ok — all {len(DIGEST['topics'])} topic headings + "
          f"{len(REQUIRED_TERMS)} key terms present in the deck")

    print(f"\nPASS — freestyle E2E clean. Output: {pptx}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
