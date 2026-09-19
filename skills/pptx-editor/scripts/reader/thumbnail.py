#!/usr/bin/env python3
"""Generate a thumbnail grid image from a PowerPoint file.

Converts every slide to a JPEG thumbnail and arranges them in a grid.
Hidden slides appear with a grey cross-hatch pattern.
Each cell is labelled with its 1-based slide number.

Usage:
    python thumbnail.py <input.pptx> [output_prefix] [--cols N]

Arguments:
    input          Path to the .pptx file
    output_prefix  Prefix for output files (default: thumbnails)
    --cols N       Number of grid columns (default: 4, max: 6)

Output:
    Single grid:    {prefix}.jpg
    Multiple grids: {prefix}-1.jpg, {prefix}-2.jpg, ...
    (A new grid starts every colsx(cols+1) slides.)

Examples:
    python thumbnail.py deck.pptx
    python thumbnail.py deck.pptx resources/deck-thumbs --cols 3
"""

import argparse
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

# Only these executables may be invoked by this script.
ALLOWED_BINARIES = {"soffice", "pdftoppm"}


def _run_allowed(cmd: list[str], **kw) -> subprocess.CompletedProcess:
    """subprocess.run wrapper that only allows a fixed set of trusted binaries."""
    if cmd[0] not in ALLOWED_BINARIES:
        raise ValueError(f"Refusing to run unapproved binary: {cmd[0]}")
    return subprocess.run(cmd, shell=False, **kw)

# -- backend detection ---------------------------------------------------------

def _detect_backend() -> str:
    """Return 'powerpoint' (Windows COM), 'soffice', or 'none'."""
    if sys.platform == "win32":
        try:
            import win32com.client  # noqa: F401
            import win32com.client as _c
            app = _c.Dispatch("PowerPoint.Application")
            app.Quit()
            return "powerpoint"
        except Exception:
            pass
    # try soffice
    r = _run_allowed(["soffice", "--version"], capture_output=True)
    if r.returncode == 0:
        return "soffice"
    return "none"

# -- tuneable constants --------------------------------------------------------
CELL_W        = 320    # px width of each thumbnail cell
DPI           = 96     # PDF -> image resolution
DEFAULT_COLS  = 4      # grid columns
MAX_COLS      = 6
JPEG_QUALITY  = 92
PAD           = 16     # px gap between cells and at canvas edges
BORDER_PX     = 2      # px border around each thumbnail
LABEL_FRAC    = 0.09   # label font size as fraction of cell width


# -- entry point ---------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    ap.add_argument("input",  help="Input .pptx file")
    ap.add_argument("prefix", nargs="?", default="thumbnails",
                    help="Output prefix (default: thumbnails)")
    ap.add_argument("--cols", type=int, default=DEFAULT_COLS,
                    help=f"Grid columns 1-{MAX_COLS} (default: {DEFAULT_COLS})")
    args = ap.parse_args()

    cols = max(1, min(args.cols, MAX_COLS))
    src  = Path(args.input)

    if not src.is_file() or src.suffix.lower() != ".pptx":
        sys.exit(f"Error: {args.input} is not a .pptx file")

    hidden   = _hidden_slides(src)      # 1-based set of hidden slide numbers
    n_total  = _total_slides(src)
    backend  = _detect_backend()
    if backend == "none":
        sys.exit(
            "Error: no rendering backend found.\n"
            "  Windows: pip install pywin32  (requires Microsoft PowerPoint)\n"
            "  Linux/Mac: install LibreOffice (soffice)"
        )

    with tempfile.TemporaryDirectory() as tmp:
        tmpdir = Path(tmp)
        pages  = _render_to_images(src, tmpdir, backend)   # visible slides only
        cells  = _assemble_cells(pages, hidden, n_total, tmpdir)
        paths  = _write_grids(cells, cols, Path(f"{args.prefix}.jpg"))

    for p in paths:
        print(p)


# -- slide metadata ------------------------------------------------------------

def _total_slides(src: Path) -> int:
    return len(Presentation(str(src)).slides)


def _hidden_slides(src: Path) -> set[int]:
    """Return 1-based numbers of slides marked hidden (show='0')."""
    prs = Presentation(str(src))
    return {
        i + 1
        for i, slide in enumerate(prs.slides)
        if slide.element.get("show") == "0"
    }


# -- PDF / image conversion ----------------------------------------------------

def _render_to_images(src: Path, outdir: Path, backend: str) -> list[Path]:
    """Convert visible slides to JPEG using the available backend."""
    if backend == "powerpoint":
        return _render_via_powerpoint(src, outdir)
    return _render_via_soffice(src, outdir)


def _render_via_powerpoint(src: Path, outdir: Path) -> list[Path]:
    """Export each slide as JPEG using PowerPoint COM automation (Windows)."""
    import win32com.client

    abs_src = str(src.resolve())
    abs_out = str(outdir.resolve())

    ppt = win32com.client.Dispatch("PowerPoint.Application")
    ppt.Visible = 1
    try:
        prs = ppt.Presentations.Open(abs_src, ReadOnly=True, Untitled=False, WithWindow=False)
        prs.Export(abs_out, "JPG")
        prs.Close()
    finally:
        ppt.Quit()

    return sorted(outdir.glob("*.jpg"))


def _render_via_soffice(src: Path, outdir: Path) -> list[Path]:
    """Convert visible slides to JPEG via PDF (soffice + pdftoppm)."""
    pdf = outdir / (src.stem + ".pdf")

    r = _run_allowed(
        ["soffice", "--headless", "--convert-to", "pdf",
         "--outdir", str(outdir), str(src)],
        capture_output=True, text=True,
    )
    if r.returncode != 0 or not pdf.exists():
        # Truncate to bound error output; never reused to build further commands/paths.
        sys.exit(f"soffice conversion failed:\n{r.stderr.strip()[:2000]}")

    r = _run_allowed(
        ["pdftoppm", "-jpeg", "-r", str(DPI), str(pdf), str(outdir / "pg")],
        capture_output=True, text=True,
    )
    if r.returncode != 0:
        sys.exit(f"pdftoppm failed:\n{r.stderr.strip()[:2000]}")

    return sorted(outdir.glob("pg-*.jpg"))


# -- cell assembly -------------------------------------------------------------

def _crosshatch(size: tuple[int, int]) -> Image.Image:
    """Create a grey crosshatch placeholder for hidden slides."""
    img  = Image.new("RGB", size, "#EBEBEB")
    draw = ImageDraw.Draw(img)
    lw   = max(3, min(size) // 80)
    draw.line([(0, 0), size],             fill="#BBBBBB", width=lw)
    draw.line([(size[0], 0), (0, size[1])], fill="#BBBBBB", width=lw)
    return img


def _assemble_cells(
    pages: list[Path],
    hidden: set[int],
    n_total: int,
    tmpdir: Path,
) -> list[tuple[Path, str]]:
    """Build the ordered list of (image_path, label) for slides 1...n_total."""
    ref_size = (1920, 1080)
    if pages:
        with Image.open(pages[0]) as im:
            ref_size = im.size

    visible_iter = iter(pages)
    cells: list[tuple[Path, str]] = []

    for num in range(1, n_total + 1):
        if num in hidden:
            ph = tmpdir / f"hidden_{num:04d}.jpg"
            _crosshatch(ref_size).save(ph, "JPEG")
            cells.append((ph, f"{num} (hidden)"))
        else:
            pg = next(visible_iter, None)
            if pg is None:
                break
            cells.append((pg, str(num)))

    return cells


# -- grid rendering ------------------------------------------------------------

def _write_grids(
    cells: list[tuple[Path, str]],
    cols: int,
    out: Path,
) -> list[str]:
    """Chunk cells into grids (max colsx(cols+1) each) and save as JPEG."""
    batch_size = cols * (cols + 1)
    chunks     = [cells[i: i + batch_size] for i in range(0, len(cells), batch_size)]
    saved: list[str] = []

    for idx, chunk in enumerate(chunks):
        img  = _render_grid(chunk, cols)
        path = (
            out
            if len(chunks) == 1
            else out.parent / f"{out.stem}-{idx + 1}{out.suffix}"
        )
        path.parent.mkdir(parents=True, exist_ok=True)
        img.save(str(path), quality=JPEG_QUALITY)
        saved.append(str(path))

    return saved


def _render_grid(cells: list[tuple[Path, str]], cols: int) -> Image.Image:
    """Compose a single grid image from a list of (image_path, label) pairs."""
    font_px = max(10, int(CELL_W * LABEL_FRAC))
    label_h = font_px + PAD          # vertical space reserved for the label

    with Image.open(cells[0][0]) as im:
        aspect = im.height / im.width
    cell_h = int(CELL_W * aspect)    # thumbnail height

    rows     = (len(cells) + cols - 1) // cols
    canvas_w = cols * CELL_W + (cols + 1) * PAD
    canvas_h = rows * (label_h + cell_h) + (rows + 1) * PAD

    canvas = Image.new("RGB", (canvas_w, canvas_h), "white")
    draw   = ImageDraw.Draw(canvas)

    try:
        font = ImageFont.load_default(size=font_px)
    except TypeError:
        font = ImageFont.load_default()

    for idx, (img_path, label) in enumerate(cells):
        col = idx % cols
        row = idx // cols

        x0    = col * CELL_W + (col + 1) * PAD
        y0    = row * (label_h + cell_h) + (row + 1) * PAD
        y_img = y0 + label_h

        # Label centered above thumbnail
        bbox = draw.textbbox((0, 0), label, font=font)
        lw   = bbox[2] - bbox[0]
        draw.text((x0 + (CELL_W - lw) // 2, y0), label, fill="#444444", font=font)

        # Thumbnail
        with Image.open(img_path) as im:
            im.thumbnail((CELL_W, cell_h), Image.Resampling.LANCZOS)
            tw, th = im.size
            tx = x0 + (CELL_W - tw) // 2
            ty = y_img + (cell_h - th) // 2
            canvas.paste(im, (tx, ty))

        # Border
        if BORDER_PX:
            draw.rectangle(
                [(tx - BORDER_PX, ty - BORDER_PX),
                 (tx + tw + BORDER_PX - 1, ty + th + BORDER_PX - 1)],
                outline="#AAAAAA", width=BORDER_PX,
            )

    return canvas


if __name__ == "__main__":
    main()
